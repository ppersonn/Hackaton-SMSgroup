"""Roda os 8 cenários (mesma semente), exporta JSON para o 3D e preenche a planilha.

Uso:
    python -m almox.cenarios                       # padrão (60 itens, 365 dias)
    python -m almox.cenarios --xlsx plano.xlsx     # lê Premissas e preenche Comparação
"""
from __future__ import annotations

import argparse
import json
import shutil
from dataclasses import asdict
from pathlib import Path

import numpy as np

from .gerador import ConfigGerador, gerar_base
from .layout import ConfigLayout, Layout
from .params import Params, carregar_premissas_xlsx, efetivos, EFEITOS
from .politicas import MinMax, SmartStock
from .sim import Mundo, Sim

CENARIOS = [
    ("Base (nada)", 0, 0, 0), ("Só RFID", 1, 0, 0), ("Só visão computacional", 0, 1, 0),
    ("Só SmartStock", 0, 0, 1), ("RFID + visão", 1, 1, 0), ("RFID + SmartStock", 1, 0, 1),
    ("Visão + SmartStock", 0, 1, 1), ("Todos", 1, 1, 1),
]
# chaves de KPI na mesma ordem em que aparecem na aba Comparação
CHAVES = ["sep_tempo", "prod", "ciclo", "dist", "busca", "receb", "contagem", "erro_sep", "retrab",
          "util", "fill", "otif", "rup_dias", "rup_itens", "diverg", "acur_loc", "fantasma", "perdas",
          "est_medio", "giro", "cobertura", "ped_compra", "emergencia", "parados", "ocupacao",
          "latencia", "rastreio", "falsos", "custo_op"]


def custos_tecnologias(xlsx=None) -> dict:
    """Custo anualizado e investimento de cada módulo (aba Módulos; padrão = exemplos da planilha)."""
    tec = {"rfid": (150000, 5, 12000), "visao": (90000, 4, 9000), "smartstock": (30000, 3, 6000)}
    if xlsx:
        from openpyxl import load_workbook
        ws = load_workbook(xlsx, data_only=True)["Módulos"]
        for k, r in zip(("rfid", "visao", "smartstock"), (5, 6, 7)):
            tec[k] = (float(ws.cell(row=r, column=3).value), float(ws.cell(row=r, column=4).value),
                      float(ws.cell(row=r, column=5).value))
    return {k: {"invest": c, "anual": round(c / max(v, 1) + m, 2)} for k, (c, v, m) in tec.items()}


def montar_mundo(n_itens=60, dias=365, semente=42, base: Params | None = None):
    base = base or Params()
    cfg = ConfigGerador(n_itens=n_itens, dias=dias, semente=semente)
    cat, dem, lt = gerar_base(cfg)
    layout = Layout(ConfigLayout(), len(cat))
    loc = layout.alocar_itens(cat)
    return Mundo(cat, dem, lt, layout, loc, base, semente), base


def escolher_dia_viz(mundo, base, minimo=20):
    """Dia com mais recebimentos no cenário base (mais coisa acontecendo para mostrar)."""
    s = Sim(mundo, efetivos(base, False, False), MinMax(), False, False).rodar()
    T = mundo.T
    cand = range(min(minimo, T - 2), T - 1)
    return max(cand, key=lambda d: (len(s.lotes_dia[d]), -abs(d - T // 2)))


def rodar_cenario(mundo, base, rfid, visao, smart, dia_viz, efeitos=None):
    P = efetivos(base, bool(rfid), bool(visao), efeitos)
    pol = (SmartStock(nivel_servico=P.nivel_servico, alfa=P.alfa, cobertura_dias=P.cobertura_dias)
           if smart else MinMax())
    return Sim(mundo, P, pol, bool(rfid), bool(visao), dia_viz=dia_viz).rodar()


def rodar_todos(n_itens=60, dias=365, semente=42, xlsx=None):
    base, efeitos, usados = Params(), None, []
    if xlsx:
        base, efeitos, usados = carregar_premissas_xlsx(xlsx)
    mundo, base = montar_mundo(n_itens, dias, semente, base)
    dia_viz = escolher_dia_viz(mundo, base)
    saidas = []
    for nome, rf, vi, sm in CENARIOS:
        s = rodar_cenario(mundo, base, rf, vi, sm, dia_viz, efeitos)
        saidas.append({"nome": nome, "rfid": rf, "visao": vi, "smartstock": sm,
                       "kpis": s.kpis(), "series": s.series(), "dia": s.dia_gravado()})
    return mundo, base, dia_viz, saidas, usados, custos_tecnologias(xlsx)


# ------------------------------------------------------------------ exportação
def exportar_json(mundo, base, dia_viz, saidas, destino: Path, tec: dict | None = None):
    destino.mkdir(parents=True, exist_ok=True)
    cat = mundo.cat
    doc = {
        "meta": {"n_itens": mundo.n, "dias": mundo.T, "dia_viz": dia_viz,
                 "turno_ini_h": base.turno_ini_h, "turno_h": base.turno_h,
                 "n_operadores": base.n_operadores, "tecnologias": tec or custos_tecnologias(), "aviso":
                 "Efeitos das tecnologias são premissas iniciais, não dados medidos."},
        "layout": mundo.layout.para_json(mundo.loc_item, cat),
        "cenarios": [{
            "nome": s["nome"], "rfid": s["rfid"], "visao": s["visao"], "smartstock": s["smartstock"],
            "kpis": {k: (round(v, 4) if isinstance(v, float) else v)
                     for k, v in s["kpis"].items() if k != "_componentes"},
            "componentes": {k: round(v, 1) for k, v in s["kpis"]["_componentes"].items()},
            "series": s["series"], "dia": s["dia"],
        } for s in saidas],
    }
    (destino / "dados.json").write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")),
                                        encoding="utf-8")
    return doc


def preencher_planilha(saidas, origem: str, destino: str):
    """Grava os KPIs nas células amarelas da aba Comparação (colunas F..M)."""
    from openpyxl import load_workbook
    shutil.copyfile(origem, destino)
    wb = load_workbook(destino)
    ws = wb["Comparação"]
    nomes = {}
    for r in range(5, 80):
        v = ws.cell(row=r, column=1).value
        if v and v not in nomes:
            nomes[v] = r                    # primeira ocorrência = linha principal (não a auxiliar)
    ROTULO = {
        "sep_tempo": "Tempo médio de separação por pedido", "prod": "Produtividade (pedidos por operador-hora)",
        "ciclo": "Tempo de ciclo do pedido (do pedido à expedição)", "dist": "Distância percorrida por pedido",
        "busca": "Tempo gasto procurando itens", "receb": "Tempo de recebimento por lote",
        "contagem": "Tempo de inventário e contagem cíclica", "erro_sep": "Taxa de erro de separação",
        "retrab": "Retrabalho e correções internas", "util": "Utilização dos operadores",
        "fill": "Nível de serviço (fill rate)", "otif": "Pedidos completos e no prazo (OTIF)",
        "rup_dias": "Dias de ruptura", "rup_itens": "Itens com ruptura",
        "diverg": "Divergência de inventário", "acur_loc": "Acurácia de localização",
        "fantasma": "Estoque fantasma", "perdas": "Perdas, avarias e extravios",
        "est_medio": "Estoque médio", "giro": "Giro anual", "cobertura": "Cobertura média de estoque",
        "ped_compra": "Pedidos de compra emitidos", "emergencia": "Compras emergenciais",
        "parados": "Itens parados (sem saída há 90 dias)", "ocupacao": "Ocupação do espaço",
        "latencia": "Latência da informação", "rastreio": "Cobertura de rastreamento",
        "falsos": "Alertas falsos", "custo_op": "Custo operacional anual (sem tecnologias)",
    }
    faltou = [k for k in CHAVES if ROTULO[k] not in nomes]
    assert not faltou, f"linhas não encontradas na planilha: {faltou}"
    for j, s in enumerate(saidas):
        col = 6 + j
        for k in CHAVES:
            ws.cell(row=nomes[ROTULO[k]], column=col, value=round(float(s["kpis"][k]), 4))
    wb.save(destino)
    return destino


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--itens", type=int, default=60)
    ap.add_argument("--dias", type=int, default=365)
    ap.add_argument("--semente", type=int, default=42)
    ap.add_argument("--xlsx", default=None, help="planilha com a aba Premissas (e Comparação a preencher)")
    ap.add_argument("--saida", default="saida")
    a = ap.parse_args()

    mundo, base, dia_viz, saidas, usados, tec = rodar_todos(a.itens, a.dias, a.semente, a.xlsx)
    out = Path(a.saida)
    exportar_json(mundo, base, dia_viz, saidas, out, tec)
    print(f"dados.json gravado em {out}/ (dia 3D = {dia_viz})")
    if usados:
        print("premissas lidas da planilha:", "; ".join(usados))
    if a.xlsx:
        dest = out / "planilha_preenchida.xlsx"
        preencher_planilha(saidas, a.xlsx, str(dest))
        print("planilha preenchida:", dest)
    print(f"{'Cenário':26s} {'sep(min)':>8s} {'prod':>6s} {'fill':>6s} {'diverg':>7s} {'custo_op':>10s}")
    for s in saidas:
        k = s["kpis"]
        print(f"{s['nome']:26s} {k['sep_tempo']:8.2f} {k['prod']:6.1f} {k['fill']:6.1%} {k['diverg']:7.1%} {k['custo_op']:10,.0f}")


if __name__ == "__main__":
    main()
